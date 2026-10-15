import io
import re
from unittest.mock import MagicMock, patch

import docx
import fitz
import numpy as np
import pytest

from worker.chunking import chunk_document
from worker.extraction import ExtractedDocument, TextSection, extract_document
from worker.worker import generate_embeddings


class WordTokenizer:
    is_fast = True
    def __call__(self, text, **kwargs):
        pairs = [(match.start(), match.end()) for match in re.finditer(r'\w+|[^\w\s]', text)]
        return {'input_ids': list(range(len(pairs))), 'offset_mapping': pairs}
    def encode(self, text, add_special_tokens=False):
        return self(text)['input_ids'] + ([0, 0] if add_special_tokens else [])
    def decode(self, ids, **kwargs):
        return 'Contexto'
    def num_special_tokens_to_add(self, **kwargs):
        return 2


def test_pdf_preserves_title_and_page():
    pdf = fitz.open()
    pdf.set_metadata({'title': 'Manual corporativo', 'author': 'Equipe'})
    for value in ('Primeira página.', 'Segunda página com contrato.'):
        pdf.new_page().insert_text((72, 72), value)
    extracted = extract_document(pdf.tobytes(), 'application/pdf', 'manual.pdf')
    pdf.close()
    assert extracted.title == 'Manual corporativo'
    assert [section.page for section in extracted.sections] == [1, 2]
    assert 'contrato' in extracted.sections[1].text
    assert extracted.metadata['author'] == 'Equipe'


def test_docx_preserves_heading_table_and_order_without_inventing_pages():
    document = docx.Document()
    document.core_properties.title = 'Contrato de serviços'
    document.add_heading('Pagamento', 1)
    document.add_paragraph('O prazo é trinta dias.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Valor'
    table.cell(0, 1).text = '100'
    document.add_heading('Rescisão', 1)
    document.add_paragraph('Aviso prévio necessário.')
    buffer = io.BytesIO()
    document.save(buffer)
    extracted = extract_document(buffer.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'contrato.docx')
    assert extracted.title == 'Contrato de serviços'
    assert [section.section for section in extracted.sections] == ['Pagamento', 'Rescisão']
    assert 'Valor\t100' in extracted.sections[0].text
    assert all(section.page is None for section in extracted.sections)


def test_html_title_hierarchy_and_no_duplicate_nested_text():
    content = b'<title>Manual</title><body><h1>Contrato</h1><h2>Pagamento</h2><p>Prazo <b>trinta</b> dias.</p><script>secret</script></body>'
    extracted = extract_document(content, 'text/html', 'manual.html')
    assert extracted.title == 'Manual'
    assert extracted.sections[-1].section == 'Contrato > Pagamento'
    assert sum(section.text.count('Prazo') for section in extracted.sections) == 1
    assert 'secret' not in ''.join(section.text for section in extracted.sections)


def test_skipped_heading_levels_preserve_sibling_hierarchy():
    document = extract_document(b'<h1>Doc</h1><h4>Deep</h4><p>First.</p><h3>Higher</h3><p>Second.</p>', 'text/html', 'doc.html')
    assert document.sections[-1].section == 'Doc > Higher'


def test_markdown_sections_preserve_code_headings_as_code():
    content = '# Manual\n\nIntrodução.\n\n## Código\n```python\n# comentário\nprint(1)\n```\n\n## Fim\nConclusão.'.encode()
    extracted = extract_document(content, 'text/markdown', 'manual.md')
    assert extracted.title == 'Manual'
    assert [section.section for section in extracted.sections] == ['Manual', 'Manual > Código', 'Manual > Fim']
    assert '# comentário' in extracted.sections[1].text


def test_normalization_and_title_override():
    extracted = extract_document('Cafe\u0301\r\ncontratu-\nal'.encode(), 'text/plain', 'texto.txt', title_override='Título')
    assert extracted.title == 'Título'
    assert extracted.sections[0].text == 'Café\ncontratual'


def test_empty_scanned_pdf_requires_ocr_or_fails():
    pdf = fitz.open()
    pdf.new_page()
    with pytest.raises(ValueError, match='OCR'):
        extract_document(pdf.tobytes(), 'application/pdf', 'scan.pdf')
    pdf.close()


@pytest.mark.parametrize('initial_text', ['', 'Page 1'])
def test_ocr_only_for_pages_without_text(initial_text):
    page = MagicMock()
    page.get_text.side_effect = [initial_text, 'Texto reconhecido.']
    page.get_images.return_value = [(1,)]
    document = MagicMock()
    document.metadata = {'title': 'Scan'}
    document.needs_pass = False
    document.__iter__.return_value = [page]
    document.__len__.return_value = 1
    document.__enter__.return_value = document
    with patch('worker.extraction.fitz.open', return_value=document):
        extracted = extract_document(b'pdf', 'application/pdf', 'scan.pdf', ocr_enabled=True)
    assert extracted.metadata['ocr_pages'] == [1]
    page.get_textpage_ocr.assert_called_once()


def test_semantic_chunks_preserve_offsets_overlap_and_source_boundaries():
    document = ExtractedDocument('Manual', [TextSection('Uma sentença curta. Outra sentença para o teste. ' * 12, page=1),
                                         TextSection('Página seguinte com outro assunto.', page=2)])
    tokenizer = WordTokenizer()
    chunks = chunk_document(document, tokenizer, max_tokens=20, overlap_tokens=3)
    assert len(chunks) > 2
    for chunk in chunks:
        source = document.sections[chunk.section_idx]
        assert chunk.content == source.text[chunk.offset_start:chunk.offset_end]
        assert chunk.page == source.page
        assert chunk.token_count <= 20
        assert len(tokenizer.encode(chunk.embedding_text, add_special_tokens=True)) <= 128
    assert chunks[1].offset_start < chunks[0].offset_end
    assert chunks[-1].page == 2


@pytest.mark.parametrize('size,overlap', [(0, 0), (10, 10), (10, -1)])
def test_invalid_semantic_chunk_configuration(size, overlap):
    with pytest.raises(ValueError):
        chunk_document(ExtractedDocument('Title', [TextSection('Some text')]), WordTokenizer(), size, overlap)


def test_embeddings_are_batched_and_dimensions_validated():
    model = MagicMock(tokenizer=WordTokenizer(), max_seq_length=128)
    model.encode.side_effect = [np.zeros((2, 384)), np.zeros((2, 384)), np.zeros((1, 384))]
    with patch('worker.worker.embedding_model', model):
        vectors = generate_embeddings(['text'] * 5, batch_size=2)
    assert len(vectors) == 5
    assert model.encode.call_count == 3
    assert [len(call.args[0]) for call in model.encode.call_args_list] == [2, 2, 1]


@pytest.mark.parametrize('values', [np.zeros((1, 128)), np.full((1, 384), np.nan)])
def test_bad_embedding_output_is_rejected(values):
    model = MagicMock(tokenizer=WordTokenizer(), max_seq_length=128)
    model.encode.return_value = values
    with patch('worker.worker.embedding_model', model), pytest.raises(ValueError, match='dimension'):
        generate_embeddings(['text'])
