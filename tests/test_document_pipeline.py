"""Real corpus: PDF/DOCX/HTML/TXT/Markdown and an image-only PDF with OCR."""

import io
import json
import math
import os
import time
import uuid
from pathlib import Path

import docx
import fitz
import httpx
import psycopg2
import pytest

URL = os.getenv("RAG_E2E_URL")
KEY = os.getenv("RAG_E2E_API_KEY")
DB = os.getenv("RAG_E2E_DATABASE_URL")
pytestmark = pytest.mark.skipif(
    not all((URL, KEY, DB)), reason="Real test stack URL/key/database required"
)


@pytest.fixture(scope="module")
def measurements():
    values = []
    yield values
    target = Path(".test-results/document-corpus.json")
    target.parent.mkdir(exist_ok=True)
    target.write_text(json.dumps(values, ensure_ascii=False, indent=2), encoding="utf-8")


def make_document(kind, marker):
    title, page, section = "Manual de contratos", None, "Pagamento"
    query = marker
    body = f"{marker} O prazo de pagamento é trinta dias."
    if kind in ("pdf", "ocr", "ocr_footer"):
        pdf = fitz.open()
        pdf.set_metadata({"title": title})
        if kind == "pdf":
            pdf.new_page().insert_text((72, 72), "Introdução e condições gerais.")
            pdf.new_page().insert_text((72, 72), body)
            page = 2
        else:
            original = fitz.open()
            original.new_page().insert_text(
                (60, 100),
                "Pagamento em trinta dias.\nDocumento digitalizado de teste.",
                fontsize=24,
            )
            pixels = original[0].get_pixmap(dpi=144)
            pdf.new_page().insert_image(pdf[-1].rect, stream=pixels.tobytes("png"))
            if kind == "ocr_footer":
                pdf[-1].insert_text((72, 800), "Page 1")
            original.close()
            page, query = 1, "Pagamento"
        data = pdf.tobytes()
        pdf.close()
        return data, "application/pdf", ".pdf", title, page, None, query
    if kind == "docx":
        document = docx.Document()
        document.core_properties.title = title
        document.add_heading(section, 1)
        document.add_paragraph(body)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Valor", "100"
        buffer = io.BytesIO()
        document.save(buffer)
        return (
            buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".docx",
            title,
            None,
            section,
            query,
        )
    if kind == "html":
        return (
            f"<title>{title}</title><body><h1>{section}</h1><p>{body}</p></body>".encode(),
            "text/html",
            ".html",
            title,
            None,
            section,
            query,
        )
    if kind == "html_skipped_heading":
        return (
            f"<title>{title}</title><body><h1>{section}</h1><h4>Detalhe</h4><p>Introdução.</p><h3>Condição</h3><p>{body}</p></body>".encode(),
            "text/html",
            ".html",
            title,
            None,
            section + " > Condição",
            query,
        )
    if kind == "md":
        return (
            f"# {title}\n\n## {section}\n\n{body}".encode(),
            "text/markdown",
            ".md",
            title,
            None,
            title + " > " + section,
            query,
        )
    if kind == "batch":
        body = f"{marker} Contrato de serviço. O prazo de pagamento é trinta dias.\n\n" * 100
    return body.encode(), "text/plain", ".txt", marker, None, None, query


@pytest.mark.parametrize(
    "kind",
    ["pdf", "docx", "html", "html_skipped_heading", "txt", "md", "batch", "ocr", "ocr_footer"],
)
def test_structured_document_pipeline(kind, measurements):
    marker = "corpus" + uuid.uuid4().hex
    data, mime, extension, title, expected_page, expected_section, query = make_document(
        kind, marker
    )
    doc_id = None
    with httpx.Client(base_url=URL, headers={"X-API-Key": KEY}, timeout=180) as client:
        try:
            upload = client.post(
                "/v1/documents",
                files={"file": (marker + extension, data, mime)},
                data={"metadata": json.dumps({"tags": [marker, "document-corpus"]})},
            )
            assert upload.status_code == 202, upload.text
            doc_id = upload.json()["id"]
            deadline = time.monotonic() + 180
            while time.monotonic() < deadline:
                response = client.get("/v1/documents/" + doc_id)
                response.raise_for_status()
                document = response.json()
                assert document["status"] != "failed", document.get("error_detail")
                if document["status"] == "indexed":
                    break
                time.sleep(1)
            else:
                pytest.fail("Document was not indexed")
            assert document["title"] == title
            assert [item["status"] for item in document["processing_history"]] == [
                "queued",
                "extracting",
                "chunking",
                "embedding",
                "indexed",
            ]
            if kind in ("ocr", "ocr_footer"):
                assert document["extraction_metadata"]["ocr_pages"] == [1]
            stats = document["extraction_metadata"]["embedding"]
            assert stats["batches"] == math.ceil(stats["chunks"] / stats["batch_size"])
            assert stats["dimension"] == 384
            assert stats["seconds"] >= 0
            if kind == "batch":
                assert stats["batches"] > 1
            measurements.append({"format": kind, **stats})
            with psycopg2.connect(DB) as connection, connection.cursor() as cursor:
                cursor.execute(
                    "SELECT id, content, offset_start, offset_end, token_count, vector_dims(embedding), chunker_version, page, section FROM chunks WHERE document_id = %s",
                    (doc_id,),
                )
                chunks = {row[0]: row for row in cursor.fetchall()}
            assert len(chunks) == stats["chunks"]
            for chunk in chunks.values():
                assert chunk[3] - chunk[2] == len(chunk[1])
                assert 0 < chunk[4] < stats["model_token_limit"]
                assert chunk[5] == 384
                assert chunk[6] == "semantic-token-chunker-v1"
            for mode in ("lexical", "vector", "hybrid"):
                result = client.post(
                    "/v1/search",
                    json={
                        "query": query,
                        "mode": mode,
                        "top_k": 5,
                        "filters": {"tags": [marker, "document-corpus"]},
                    },
                )
                assert result.status_code == 200, result.text
                hits = result.json()["results"]
                assert hits and all(hit["document_id"] == doc_id for hit in hits)
                assert any(
                    hit["citation"]["page"] == expected_page
                    and hit["citation"]["section"] == expected_section
                    for hit in hits
                )
                for hit in hits:
                    citation = hit["citation"]
                    assert citation["document_id"] == doc_id
                    assert citation["filename"] == marker + extension
                    assert citation["title"] == title
                    assert citation["page"] == chunks[hit["chunk_id"]][7]
                    assert citation["section"] == chunks[hit["chunk_id"]][8]
                    assert citation["excerpt"] == chunks[hit["chunk_id"]][1][:200]
                    assert citation["offset_end"] - citation["offset_start"] == len(
                        citation["excerpt"]
                    )
            empty = client.post(
                "/v1/search",
                json={
                    "query": query,
                    "mode": "lexical",
                    "filters": {"tags": [marker, "does-not-exist"]},
                },
            )
            assert empty.status_code == 200 and empty.json()["results"] == []
        finally:
            if doc_id:
                cleanup = client.delete("/v1/documents/" + doc_id)
                assert cleanup.status_code in (204, 404), cleanup.text
